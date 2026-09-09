import openpyxl
from pptx import Presentation
from pptx.util import Emu, Pt

from app.core.errors import PDFError
from app.core.pdf_ops import open_pdf

_EMU_PER_POINT = 12700

# Maps this codebase's three base-14-style families (the same three
# app/core/pdf_ops.py's _base14_alias already recognizes for user-chosen
# fonts) to real font name strings python-pptx's run.font.name expects.
# Distinct from _base14_alias: that function turns an ALREADY-CHOSEN family
# into a fitz-internal alias code ("helv", "hebo", ...) for page.insert_text()
# — not a real font name, and not something that can guess a family from a
# raw detected font string. This module needs the latter, which didn't
# already exist anywhere in the codebase.
_PPTX_FONT_NAMES = {
    "helvetica": "Arial",
    "times": "Times New Roman",
    "courier": "Courier New",
}


def _closest_base14_family(font_name: str) -> str:
    name = font_name.lower()
    if "courier" in name or "mono" in name:
        return "courier"
    if any(hint in name for hint in ("times", "serif", "georgia", "garamond", "cambria", "minion")):
        return "times"
    return "helvetica"


def pdf_to_pptx(input_path: str, output_path: str) -> None:
    doc = open_pdf(input_path)
    try:
        prs = Presentation()
        slide_size_set = False
        for page in doc:
            if not slide_size_set:
                # .pptx's slide size is presentation-wide, not per-slide (verified
                # empirically: re-setting it before a later add_slide() call
                # silently resizes the whole deck once saved and reopened) — set
                # it once, from the first page. Correct and lossless for the
                # common case of a uniform-page-size document; a later page of a
                # different size still gets text-accurate, absolutely-positioned
                # boxes, just against a canvas sized to page 1.
                prs.slide_width = Emu(round(page.rect.width * _EMU_PER_POINT))
                prs.slide_height = Emu(round(page.rect.height * _EMU_PER_POINT))
                slide_size_set = True

            slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
            blocks = page.get_text("dict")["blocks"]
            for block in blocks:
                if "lines" not in block:
                    continue  # an image block, not text
                x0, y0, x1, y1 = block["bbox"]
                textbox = slide.shapes.add_textbox(
                    Emu(round(x0 * _EMU_PER_POINT)),
                    Emu(round(y0 * _EMU_PER_POINT)),
                    Emu(round((x1 - x0) * _EMU_PER_POINT)),
                    Emu(round((y1 - y0) * _EMU_PER_POINT)),
                )
                text_frame = textbox.text_frame
                text_frame.word_wrap = True
                for line_index, line in enumerate(block["lines"]):
                    paragraph = text_frame.paragraphs[0] if line_index == 0 else text_frame.add_paragraph()
                    for span in line["spans"]:
                        run = paragraph.add_run()
                        run.text = span["text"]
                        run.font.size = Pt(span["size"])
                        flags = span["flags"]
                        run.font.bold = bool(flags & 16)
                        run.font.italic = bool(flags & 2)
                        run.font.name = _PPTX_FONT_NAMES[_closest_base14_family(span["font"])]
        prs.save(output_path)
    finally:
        doc.close()


def pdf_to_xlsx(input_path: str, output_path: str) -> None:
    doc = open_pdf(input_path)
    try:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)  # default blank sheet — replaced by real ones below
        table_count = 0
        for page_index, page in enumerate(doc, start=1):
            tables = page.find_tables()
            for table_index, table in enumerate(tables.tables, start=1):
                table_count += 1
                sheet = workbook.create_sheet(f"Page{page_index}_Table{table_index}")
                for row in table.extract():
                    sheet.append(row)
        if table_count == 0:
            raise PDFError("No tables found in this document.")
        workbook.save(output_path)
    finally:
        doc.close()
