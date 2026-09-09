import fitz
import pytest
from pptx import Presentation

from app.core.errors import PDFError
from app.core.pdf_to_office import pdf_to_pptx


def _make_text_pdf(path):
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 100), "Heading Text", fontname="hebo", fontsize=24)
    page.insert_text((72, 200), "Body paragraph in times.", fontname="tiro", fontsize=12)
    doc.save(str(path))
    doc.close()


def test_pdf_to_pptx_creates_positioned_textboxes(tmp_path):
    input_path = tmp_path / "input.pdf"
    _make_text_pdf(input_path)

    output_path = tmp_path / "output.pptx"
    pdf_to_pptx(str(input_path), str(output_path))

    assert output_path.exists()
    prs = Presentation(str(output_path))
    assert len(prs.slides) == 1
    shapes = list(prs.slides[0].shapes)
    assert len(shapes) == 2

    # Exact EMU values verified empirically against this exact fixture this
    # session (1pt = 12700 EMU, round(value_in_points * 12700)) — not
    # computed by hand.
    heading_shape = shapes[0]
    assert (heading_shape.left, heading_shape.top, heading_shape.width, heading_shape.height) == (
        914400,
        943864,
        1913839,
        419710,
    )
    heading_text = "".join(run.text for para in heading_shape.text_frame.paragraphs for run in para.runs)
    assert heading_text == "Heading Text"

    body_shape = shapes[1]
    assert (body_shape.left, body_shape.top, body_shape.width, body_shape.height) == (
        914400,
        2379523,
        1540764,
        203302,
    )
    body_text = "".join(run.text for para in body_shape.text_frame.paragraphs for run in para.runs)
    assert body_text == "Body paragraph in times."


def test_pdf_to_pptx_preserves_font_size_bold_and_family(tmp_path):
    input_path = tmp_path / "input.pdf"
    _make_text_pdf(input_path)

    output_path = tmp_path / "output.pptx"
    pdf_to_pptx(str(input_path), str(output_path))

    prs = Presentation(str(output_path))
    shapes = list(prs.slides[0].shapes)

    heading_run = shapes[0].text_frame.paragraphs[0].runs[0]
    # "hebo" (bold Helvetica) -> PyMuPDF reports font "Helvetica-Bold", flags
    # bit 4 (16) set for bold, bit 1 (2) clear for italic — verified this
    # session against this exact fixture.
    assert int(heading_run.font.size) == 304800  # Pt(24)
    assert heading_run.font.bold is True
    assert heading_run.font.italic is False
    assert heading_run.font.name == "Arial"

    body_run = shapes[1].text_frame.paragraphs[0].runs[0]
    # "tiro" (plain Times) -> PyMuPDF reports font "Times-Roman", flags bit 2
    # (4, serifed) set, bold/italic bits clear.
    assert int(body_run.font.size) == 152400  # Pt(12)
    assert body_run.font.bold is False
    assert body_run.font.italic is False
    assert body_run.font.name == "Times New Roman"


def test_pdf_to_pptx_matches_slide_size_to_page_size(tmp_path):
    input_path = tmp_path / "input.pdf"
    _make_text_pdf(input_path)

    output_path = tmp_path / "output.pptx"
    pdf_to_pptx(str(input_path), str(output_path))

    prs = Presentation(str(output_path))
    # 612pt x 792pt (US Letter) -> round(612*12700), round(792*12700) —
    # verified empirically this session.
    assert (prs.slide_width, prs.slide_height) == (7772400, 10058400)


def test_pdf_to_pptx_skips_image_blocks(tmp_path):
    doc = fitz.open()
    page = doc.new_page(width=200, height=200)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10))
    pix.set_rect(pix.irect, (255, 0, 0))
    page.insert_image(fitz.Rect(10, 10, 50, 50), pixmap=pix)
    page.insert_text((10, 100), "Some text", fontsize=12)
    input_path = tmp_path / "input.pdf"
    doc.save(str(input_path))
    doc.close()

    output_path = tmp_path / "output.pptx"
    pdf_to_pptx(str(input_path), str(output_path))

    prs = Presentation(str(output_path))
    shapes = list(prs.slides[0].shapes)
    # Only the text block becomes a text box — the inserted image produces a
    # block with no "lines" key and is skipped entirely (confirmed
    # empirically this session: image blocks are type==1 with no "lines"
    # key; text blocks are type==0 with a "lines" key).
    assert len(shapes) == 1
    assert shapes[0].text_frame.paragraphs[0].runs[0].text == "Some text"


def test_pdf_to_pptx_raises_for_missing_file(tmp_path):
    with pytest.raises(PDFError):
        pdf_to_pptx(str(tmp_path / "does_not_exist.pdf"), str(tmp_path / "output.pptx"))
